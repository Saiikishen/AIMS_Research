#!/usr/bin/env python3
"""
Convert "Pre surgical case Conference" notes stored in .eml files into an
organised Excel workbook.

  * One .eml can contain several cases -> one row per case on the "Cases" sheet
  * Every section (Clinical History, Video EEG, MRI, Wada, Plan ...) gets its own column
  * A second sheet, "Sections", shows the same data in a long, easy-to-read layout

Usage:
    python eml_cases_to_excel.py "C:/path/to/folder"            # all .eml files, sub-folders included
    python eml_cases_to_excel.py "C:/path/to/folder" -o out.xlsx # choose the output file
    python eml_cases_to_excel.py "C:/path/to/folder" --per-file  # one .xlsx per .eml
    python eml_cases_to_excel.py                                  # asks you for the folder path

Requires: pip install openpyxl pandas

HOW TO ADD A NEW SECTION
    Add a line to HEADINGS below: ("Column name in Excel", r"regex of the heading text").
    Headings are only recognised at the START of a line and must be followed by a
    colon. Anything the script does not recognise stays inside the section above it,
    so no text is ever lost.
"""

import argparse
import re
from datetime import datetime
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Heading definitions: (Excel column, regex for the heading label)
# Several headings may point to the same column; their text is then combined.
# --------------------------------------------------------------------------
HEADINGS = [
    # ---- patient / case details -------------------------------------------
    ("Name", r"Name"),
    ("MRD", r"MRD(?:[ \t]*(?:No|Number)\.?)?"),
    ("Mobile", r"Mob(?:ile)?[ \t]*(?:No|Number)\.?"),  # colon optional
    ("DOB", r"DOB|Date of Birth"),
    ("Age", r"Age"),
    ("Sex", r"Sex|Gender"),
    ("Handedness", r"Edinburgh Handedness|Handedness"),
    ("Place", r"Place"),
    ("Referring Doctor", r"Referring doctor"),
    ("Referred By", r"Referred by"),
    ("Conference Date", r"Conference Date"),
    ("Reason for Presentation", r"Reason for presentation"),
    # ---- clinical ----------------------------------------------------------
    ("Clinical History", r"Clinical History"),
    ("Age at Onset", r"Age at onset"),
    ("Duration of Illness", r"(?:Total[ \t]+)?Duration of illness"),
    ("Seizure Semiology", r"Seizure semiology"),
    ("Seizure Frequency", r"(?:Seizure[ \t]+|Current[ \t]+)?Frequency"),
    ("Current Seizure Semiology", r"Current seizure semiology"),
    ("Status Epilepticus", r"Status epilepticus"),
    ("Birth and Development", r"Birth and Development"),
    ("Febrile Seizures", r"Febrile seizures?"),
    ("Family History", r"Family history"),
    ("Social History", r"Social History"),
    ("Previous Medications", r"Previous Meds|Medications?[ \t]+history"),
    ("Current Medications", r"Current Meds"),
    ("Examination", r"Examination"),
    # ---- investigations ----------------------------------------------------
    ("Scalp EEGs", r"Scalp EEGs?"),
    ("Neuropsychological Assessment", r"Neuropsychological Assessment|Neuropsychology"),
    ("Neuro-ophthalmology", r"Neuro-?ophthalmology|HFA|CVI"),
    ("DTI", r"DTI(?:[ \t]+results)?"),
    ("Video EEG", r"Video EEG"),
    ("Interictal EEG", r"Interictal"),
    ("Ictal EEG", r"Ictal(?![ \t]*ESI)"),
    ("Ictal ESI", r"Ictal[ \t]*ESI"),
    ("Electroclinical Correlation", r"Electroclinical Correlation"),
    ("Localization", r"Localization"),
    ("MRI", r"MRI"),
    ("Hippocampal Volumetry", r"Hippocampal volumetry|HV"),
    ("Morphometric Analysis", r"Morphometric analysis"),
    ("PET", r"PET"),
    ("SISCOM", r"SISCOM"),
    ("Wada", r"Wada"),
    # ---- decision ----------------------------------------------------------
    ("Pre-surgical Decision & Plan",
     r"Pre[ \t-]*surgical[ \t]+Decision(?:[ \t]*(?:&|and)[ \t]*Plan)?"),
]

COLON_OPTIONAL = {"Mobile"}  # "Mob NO.9734439832" has no colon

HEADER_COLS = [
    "Name", "MRD", "Conference Date", "Mobile", "DOB", "Age", "Sex", "Handedness",
    "Place", "Referring Doctor", "Referred By", "Reason for Presentation",
]
SECTION_COLS = []
for _col, _ in HEADINGS:
    if _col not in HEADER_COLS and _col not in SECTION_COLS:
        SECTION_COLS.append(_col)

META_COLS = [
    "Images Referenced", "Source File", "Email Subject", "Email Date",
    "Original Sender", "Original Sent", "Case #", "Full Case Text",
]
DATE_COLS = {"DOB", "Conference Date"}

# Fields that hold a single short value; extra lines below them are really the
# seizure-semiology narrative (e.g. "Total Duration of illness: 14 years" followed
# by "Semiology - Initial 2013 ...").
SINGLE_LINE = {"Age at Onset", "Duration of Illness"}
SPILL_INTO = "Seizure Semiology"

CASE_SPLIT = re.compile(r"pre[\s-]*surgical\s+case\s+conference", re.I)
IMG_LINE = re.compile(
    r"^[ \t]*\[?[ \t]*Pasted Graphic[^\n\]]*?\.(?:png|jpe?g|gif|bmp)[ \t]*\]?[ \t]*$",
    re.I | re.M,
)
ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MAX_CELL = 32000


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def normalize(text: str) -> str:
    """Tidy line endings / whitespace and squash the blank line between every line."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = ILLEGAL_XML.sub("", text)
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()


def html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?>.*?</\1>", "", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>", "\n", html)
    return re.sub(r"<[^>]+>", "", html)


def get_body(msg) -> str:
    plain = msg.get_body(preferencelist=("plain",))
    if plain is not None:
        return plain.get_content()
    html = msg.get_body(preferencelist=("html",))
    return html_to_text(html.get_content()) if html is not None else ""


def parse_date(value: str):
    value = value.strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return value  # keep the original text if it isn't a recognisable date


def cell(text) -> str:
    return ILLEGAL_XML.sub("", str(text))[:MAX_CELL] if text else ""


# --------------------------------------------------------------------------
# Parsing one case
# --------------------------------------------------------------------------
def parse_case(raw: str) -> dict:
    text = normalize(raw)

    # 1. find every heading in the case
    hits = []
    for col, pattern in HEADINGS:
        colon = r"[ \t]*:?" if col in COLON_OPTIONAL else r"[ \t]*:"
        rx = re.compile(rf"^[ \t]*(?P<label>{pattern}){colon}", re.I | re.M)
        for m in rx.finditer(text):
            hits.append((m.start(), m.end(), col, m.group("label").strip()))
    hits.sort(key=lambda h: (h[0], -(h[1] - h[0])))

    unique, last_start = [], -1
    for h in hits:
        if h[0] != last_start:
            unique.append(h)
            last_start = h[0]

    # 2. slice the text between consecutive headings
    pieces, images = {}, []
    for i, (_, end, col, label) in enumerate(unique):
        stop = unique[i + 1][0] if i + 1 < len(unique) else len(text)
        content = text[end:stop]
        for img in IMG_LINE.findall(content):
            images.append(f"{label}: {img.strip(' []')}")
        content = normalize(IMG_LINE.sub("", content))
        if col in SINGLE_LINE and "\n" in content:
            content, spill = content.split("\n", 1)
            pieces.setdefault(SPILL_INTO, []).append((None, normalize(spill)))
        if content:
            pieces.setdefault(col, []).append((label, content))

    # 3. build the row (if a column got several pieces, keep each piece's label)
    row = {}
    for col, items in pieces.items():
        if len(items) == 1:
            value = items[0][1]
        else:
            value = "\n".join(f"{lab}: {txt}" if lab else txt for lab, txt in items)
        if col in HEADER_COLS:
            value = value.replace("\n", " ").strip()
        row[col] = value

    for col in DATE_COLS:
        if col in row:
            row[col] = parse_date(row[col])

    row["Images Referenced"] = "\n".join(images)
    row["Full Case Text"] = text
    return row


# --------------------------------------------------------------------------
# Parsing a whole .eml
# --------------------------------------------------------------------------
def parse_eml(path: Path, base: Path | None = None) -> list[dict]:
    with open(path, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)

    body = get_body(msg)
    chunks = CASE_SPLIT.split(body)
    preamble, case_chunks = chunks[0], chunks[1:]

    if not case_chunks:  # no "Pre surgical case Conference" marker found
        # no "Pre surgical case Conference" marker: still try the whole body as one case
        preamble, case_chunks = "", [body]

    sender = re.search(r"(?im)^From:[ \t]*(.+)$", preamble)
    sent = re.search(r"(?im)^Sent:[ \t]*(.+)$", preamble)

    try:
        email_date = parsedate_to_datetime(msg["Date"]).replace(tzinfo=None)
    except Exception:
        email_date = None

    rows = []
    for n, chunk in enumerate(case_chunks, start=1):
        if not chunk.strip():
            continue
        row = parse_case(chunk)
        if not row.get("Name"):  # not a recognisable case note (e.g. an unrelated email)
            continue
        row.update({
            "Case #": n,
            "Source File": str(path.relative_to(base)) if base else path.name,
            "Email Subject": cell(str(msg["Subject"] or "").replace("\xa0", " ").strip()),
            "Email Date": email_date,
            "Original Sender": sender.group(1).strip() if sender else "",
            "Original Sent": sent.group(1).strip() if sent else "",
        })
        rows.append(row)
    return rows


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------
def build_frames(rows: list[dict]):
    cases = pd.DataFrame(rows)

    ordered = HEADER_COLS + SECTION_COLS + META_COLS
    for col in ordered:
        if col not in cases.columns:
            cases[col] = ""
    cases = cases[ordered].fillna("")

    # drop section columns that are empty for every case
    empty = [c for c in SECTION_COLS if (cases[c] == "").all()]
    cases = cases.drop(columns=empty)

    long_rows = []
    for _, r in cases.iterrows():
        for col in SECTION_COLS:
            if col in cases.columns and r[col] != "":
                long_rows.append({
                    "Name": r["Name"], "MRD": r["MRD"],
                    "Conference Date": r["Conference Date"],
                    "Section": col, "Content": r[col],
                })
    sections = pd.DataFrame(long_rows, columns=["Name", "MRD", "Conference Date", "Section", "Content"])
    return cases, sections


def style_sheet(ws, df: pd.DataFrame, widths: dict, default_width=55, freeze="B2"):
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    body_font = Font(name="Arial", size=10)

    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for idx, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, default_width)

    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body_font
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if hasattr(c.value, "year"):
                c.number_format = "dd-mm-yyyy"

    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def write_excel(rows: list[dict], output: Path):
    cases, sections = build_frames(rows)

    narrow = {c: 16 for c in HEADER_COLS}
    narrow.update({"Name": 24, "Reason for Presentation": 35, "Referring Doctor": 22,
                   "Case #": 8, "Source File": 28, "Email Subject": 24, "Email Date": 18,
                   "Original Sender": 32, "Original Sent": 22, "Images Referenced": 30,
                   "Full Case Text": 90, "Age at Onset": 16, "Duration of Illness": 16,
                   "Status Epilepticus": 20, "Febrile Seizures": 20})

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        cases.to_excel(writer, sheet_name="Cases", index=False)
        sections.to_excel(writer, sheet_name="Sections", index=False)
        style_sheet(writer.sheets["Cases"], cases, narrow)
        style_sheet(
            writer.sheets["Sections"], sections,
            {"Name": 24, "MRD": 12, "Conference Date": 16, "Section": 30, "Content": 110},
            freeze="A2",
        )


# --------------------------------------------------------------------------
# Command line
# --------------------------------------------------------------------------
def find_eml_files(folder: Path) -> list[Path]:
    """Every .eml under `folder`, including all sub-folders (any capitalisation)."""
    return sorted(
        p for p in folder.rglob("*")
        if p.is_file()
        and p.suffix.lower() == ".eml"
        and not p.name.startswith(("._", "~$"))  # skip macOS / Office junk files
    )


def main():
    ap = argparse.ArgumentParser(
        description="Convert case-conference .eml notes (a folder, searched recursively) to Excel."
    )
    ap.add_argument("input", nargs="?", help="folder containing .eml files (or a single .eml file)")
    ap.add_argument("-o", "--output",
                    help="output .xlsx for the combined workbook (default: <folder>/all_cases.xlsx)")
    ap.add_argument("--per-file", action="store_true",
                    help="write one .xlsx next to each .eml instead of one combined workbook")
    args = ap.parse_args()

    raw = args.input or input("Path to the folder with your .eml files: ")
    src = Path(raw.strip().strip('"').strip("'")).expanduser()

    if src.is_file():
        files, base = [src], src.parent
    elif src.is_dir():
        files, base = find_eml_files(src), src
    else:
        raise SystemExit(f"Path not found: {src}")
    if not files:
        raise SystemExit(f"No .eml files found in {src} (sub-folders included).")

    print(f"Found {len(files)} .eml file(s) under {src}\n")

    all_rows, failed, empty = [], [], []
    for i, f in enumerate(files, start=1):
        rel = f.relative_to(base)
        try:
            found = parse_eml(f, base)
        except Exception as exc:  # keep going if one file is broken
            failed.append((rel, exc))
            print(f"[{i}/{len(files)}] {rel}: FAILED ({exc})")
            continue
        if not found:
            empty.append(rel)
        print(f"[{i}/{len(files)}] {rel}: {len(found)} case(s)")

        if args.per_file and found:
            out = f.with_suffix(".xlsx")
            write_excel(found, out)
        all_rows.extend(found)

    if not all_rows:
        raise SystemExit("\nNo cases could be parsed.")

    if args.per_file:
        print(f"\nWrote {len(all_rows)} case(s) into individual .xlsx files next to each .eml")
    else:
        out = Path(args.output) if args.output else (
            src.with_name(src.stem + "_cases.xlsx") if src.is_file() else src / "all_cases.xlsx"
        )
        try:
            write_excel(all_rows, out)
        except PermissionError:
            raise SystemExit(f"\nCannot write {out} - is it open in Excel? Close it and run again.")
        print(f"\nWrote {len(all_rows)} case(s) from {len(files)} file(s) to {out}")

    if failed:
        print(f"\n{len(failed)} file(s) failed:")
        for rel, exc in failed:
            print(f"  - {rel}: {exc}")
    if empty:
        print(f"\n{len(empty)} file(s) contained no case notes (skipped):")
        for rel in empty:
            print(f"  - {rel}")


if __name__ == "__main__":
    main()
